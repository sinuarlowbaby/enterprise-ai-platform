"""
Unit tests for Document Parsers (PyMuPDF, DOCX, Markdown, Text, and Factory).
"""

import io
import fitz
import docx
import pytest

from knowledge.parsers.base import BlockType
from knowledge.parsers.docx import DocxParser
from knowledge.parsers.factory import ParserFactory
from knowledge.parsers.markdown import MarkdownParser
from knowledge.parsers.pdf import PyMuPDFParser
from knowledge.parsers.text import TextParser


def _create_sample_pdf_bytes() -> bytes:
    """Generates a multi-page PDF in memory with known font sizes and layout."""
    doc = fitz.open()

    # Page 1: Title and Section 1
    page1 = doc.new_page(width=595, height=842)
    # Title (font size 24)
    page1.insert_text((50, 80), "Enterprise AI Architecture", fontsize=24)
    # H1 (font size 18)
    page1.insert_text((50, 140), "1. Executive Summary", fontsize=18)
    # Body paragraph (font size 11)
    page1.insert_text(
        (50, 180),
        "This platform provides isolated multi-tenant hybrid retrieval and agent workflows.",
        fontsize=11,
    )
    # H2 (font size 14)
    page1.insert_text((50, 230), "1.1 Key Objectives", fontsize=14)
    page1.insert_text((50, 260), "Ensure zero cross-tenant data leakage.", fontsize=11)

    # Page 2: Section 2
    page2 = doc.new_page(width=595, height=842)
    # H1 (font size 18)
    page2.insert_text((50, 80), "2. Database Specifications", fontsize=18)
    page2.insert_text((50, 120), "PostgreSQL 16 with pgvector extension is used for storage.", fontsize=11)

    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def _create_sample_docx_bytes() -> bytes:
    """Generates a DOCX document in memory with headings, paragraphs, and a table."""
    doc = docx.Document()
    doc.core_properties.title = "Financial Report Q3"
    doc.core_properties.author = "Finance Team"

    doc.add_heading("Financial Report Q3", level=0)  # Title
    doc.add_heading("Revenue Breakdown", level=1)
    doc.add_paragraph("Quarterly gross revenue exceeded projections by 14%.")

    doc.add_heading("Regional Performance", level=2)
    doc.add_paragraph("North America accounted for 60% of total revenue.")

    # Table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "Revenue"
    table.rows[1].cells[0].text = "NA"
    table.rows[1].cells[1].text = "$12M"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_pymupdf_parser_layout_extraction():
    """Verify PyMuPDF font size analysis, heading detection, and page tracking."""
    pdf_bytes = _create_sample_pdf_bytes()
    parser = PyMuPDFParser()
    doc = parser.parse(pdf_bytes, filename="arch_spec.pdf")

    assert doc.source_filename == "arch_spec.pdf"
    assert doc.file_type == "application/pdf"
    assert doc.total_pages == 2
    assert doc.title == "Enterprise AI Architecture"

    # Check extracted blocks
    block_texts = [b.text for b in doc.blocks]
    assert any("Enterprise AI Architecture" in t for t in block_texts)
    assert any("1. Executive Summary" in t for t in block_texts)
    assert any("2. Database Specifications" in t for t in block_texts)

    # Check page numbering and block types
    p1_blocks = [b for b in doc.blocks if b.page_number == 1]
    p2_blocks = [b for b in doc.blocks if b.page_number == 2]
    assert len(p1_blocks) > 0
    assert len(p2_blocks) > 0

    # Verify section path breadcrumb on Page 2 block
    p2_content_block = next(b for b in p2_blocks if "pgvector" in b.text)
    assert "2. Database Specifications" in p2_content_block.section_path


def test_docx_parser_styles_and_tables():
    """Verify DOCX parser heading hierarchy and markdown table generation."""
    docx_bytes = _create_sample_docx_bytes()
    parser = DocxParser()
    doc = parser.parse(docx_bytes, filename="q3_report.docx")

    assert doc.title == "Financial Report Q3"
    assert doc.metadata["author"] == "Finance Team"

    # Verify table block extraction
    table_blocks = [b for b in doc.blocks if b.block_type == BlockType.TABLE]
    assert len(table_blocks) == 1
    assert "| Region | Revenue |" in table_blocks[0].text
    assert "| NA | $12M |" in table_blocks[0].text

    # Verify heading breadcrumbs
    sub_para = next(b for b in doc.blocks if "North America" in b.text)
    assert "Revenue Breakdown" in sub_para.section_path
    assert "Regional Performance" in sub_para.section_path


def test_markdown_parser_frontmatter_and_headings():
    """Verify Markdown frontmatter parsing, ATX headings, and code blocks."""
    md_content = """---
title: System Overview
author: Platform Architect
version: 1.0
---

# System Overview

High-level architecture documentation.

## Core Modules

### Ingestion Service

The ingestion service parses heterogeneous documents.

```python
def ingest_file(path: str):
    return parse(path)
```

| Component | Status |
| --- | --- |
| Parser | Active |
| Splitter | Active |
"""
    parser = MarkdownParser()
    doc = parser.parse(md_content, filename="overview.md")

    assert doc.title == "System Overview"
    assert doc.metadata.get("author") == "Platform Architect"
    assert doc.metadata.get("version") == "1.0"

    # Check code block preservation
    code_blocks = [b for b in doc.blocks if b.block_type == BlockType.CODE]
    assert len(code_blocks) == 1
    assert "def ingest_file" in code_blocks[0].text
    assert code_blocks[0].metadata.get("lang") == "python"

    # Check table block
    table_blocks = [b for b in doc.blocks if b.block_type == BlockType.TABLE]
    assert len(table_blocks) == 1
    assert "Parser | Active" in table_blocks[0].text

    # Check section path on code block
    assert "Core Modules" in code_blocks[0].section_path
    assert "Ingestion Service" in code_blocks[0].section_path


def test_text_parser_heuristics():
    """Verify plain text title and uppercase heading heuristics."""
    txt_content = """PLATFORM SPECIFICATIONS

1. INTRODUCTION
This is the introduction section.

2. SECURITY GUARDRAILS
Presidio scans all input prompts for PII.
"""
    parser = TextParser()
    doc = parser.parse(txt_content, filename="spec.txt")

    assert doc.title == "PLATFORM SPECIFICATIONS"
    headings = doc.section_headings
    assert any("INTRODUCTION" in h for h in headings)
    assert any("SECURITY GUARDRAILS" in h for h in headings)


def test_parser_factory_resolution():
    """Verify ParserFactory returns correct parser for extensions and MIME types."""
    assert isinstance(ParserFactory.get_parser("document.pdf"), PyMuPDFParser)
    assert isinstance(ParserFactory.get_parser("application/pdf"), PyMuPDFParser)
    assert isinstance(ParserFactory.get_parser("notes.docx"), DocxParser)
    assert isinstance(ParserFactory.get_parser("guide.md"), MarkdownParser)
    assert isinstance(ParserFactory.get_parser("data.txt"), TextParser)
    assert isinstance(ParserFactory.get_parser("unknown.xyz"), TextParser)
