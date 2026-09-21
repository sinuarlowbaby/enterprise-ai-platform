"""
DOCX Parser with Style & Table Extraction.

Extracts headings, paragraphs, and structured tables from Word (.docx) documents
using python-docx.
"""

import io
from pathlib import Path
from typing import Any, BinaryIO
import docx

from knowledge.parsers.base import BaseParser, BlockType, DocumentBlock, ParsedDocument


class DocxParser(BaseParser):
    """
    Parser for Microsoft Word (.docx) documents.

    Features:
    - Reads paragraph styles (Title, Heading 1..6, Normal, List).
    - Extracts tables into structured tabular text blocks.
    - Maintains section breadcrumb paths.
    - Captures core document properties (title, author, created).
    """

    def parse(
        self,
        content: str | Path | bytes | BinaryIO,
        filename: str = "document.docx",
        **kwargs: Any,
    ) -> ParsedDocument:
        doc = self._open_document(content)

        core_props = doc.core_properties
        title = core_props.title if core_props.title and core_props.title.strip() else None

        blocks: list[DocumentBlock] = []
        heading_stack: list[tuple[int, str]] = []

        # Iterate through paragraphs in order
        for p_idx, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name.lower() if para.style else "normal"
            block_type = BlockType.PARAGRAPH
            heading_level: int | None = None

            if "title" in style_name:
                block_type = BlockType.TITLE
                heading_level = 1
                if not title:
                    title = text
            elif "heading 1" in style_name:
                block_type = BlockType.HEADING
                heading_level = 1
            elif "heading 2" in style_name:
                block_type = BlockType.HEADING
                heading_level = 2
            elif "heading 3" in style_name:
                block_type = BlockType.HEADING
                heading_level = 3
            elif "heading 4" in style_name or "heading 5" in style_name or "heading 6" in style_name:
                block_type = BlockType.HEADING
                heading_level = 4
            elif "list" in style_name or text.startswith(("- ", "• ", "* ")):
                block_type = BlockType.LIST

            # Manage section hierarchy
            if heading_level is not None:
                while heading_stack and heading_stack[-1][0] >= heading_level:
                    heading_stack.pop()
                heading_stack.append((heading_level, text))

            current_section_path = [h[1] for h in heading_stack]

            block = DocumentBlock(
                text=text,
                block_type=block_type,
                page_number=1,  # DOCX does not natively store fixed page breaks
                heading_level=heading_level,
                section_path=list(current_section_path),
                metadata={
                    "style": para.style.name if para.style else "Normal",
                    "paragraph_index": p_idx,
                },
            )
            blocks.append(block)

        # Iterate through tables and format them as Markdown table blocks
        for t_idx, table in enumerate(doc.tables):
            table_md = self._table_to_markdown(table)
            if table_md.strip():
                current_section_path = [h[1] for h in heading_stack]
                table_block = DocumentBlock(
                    text=table_md,
                    block_type=BlockType.TABLE,
                    page_number=1,
                    section_path=list(current_section_path),
                    metadata={
                        "table_index": t_idx,
                        "row_count": len(table.rows),
                        "col_count": len(table.columns) if table.rows else 0,
                    },
                )
                blocks.append(table_block)

        return ParsedDocument(
            source_filename=filename,
            file_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            title=title,
            total_pages=1,
            blocks=blocks,
            metadata={
                "author": core_props.author,
                "created": str(core_props.created) if core_props.created else None,
                "modified": str(core_props.modified) if core_props.modified else None,
                "paragraph_count": len(doc.paragraphs),
                "table_count": len(doc.tables),
            },
        )

    def _open_document(self, content: str | Path | bytes | BinaryIO) -> docx.Document:
        """Opens docx document from filepath, bytes, or stream."""
        if isinstance(content, (str, Path)):
            return docx.Document(str(content))
        elif isinstance(content, bytes):
            return docx.Document(io.BytesIO(content))
        elif hasattr(content, "read"):
            data = content.read()
            return docx.Document(io.BytesIO(data))
        raise ValueError(f"Unsupported content type for DOCX parser: {type(content)}")

    def _table_to_markdown(self, table: Any) -> str:
        """Renders python-docx Table object as a GitHub Flavored Markdown table."""
        if not table.rows:
            return ""

        rows_data: list[list[str]] = []
        for row in table.rows:
            row_cells = [cell.text.replace("\n", " ").strip() for cell in row.cells]
            rows_data.append(row_cells)

        if not rows_data:
            return ""

        headers = rows_data[0]
        col_count = len(headers)
        lines: list[str] = [
            "| " + " | ".join(headers) + " |",
            "| " + " | ".join(["---"] * col_count) + " |",
        ]

        for data_row in rows_data[1:]:
            padded_row = data_row + [""] * (col_count - len(data_row))
            lines.append("| " + " | ".join(padded_row[:col_count]) + " |")

        return "\n".join(lines)
